"""Excel 导出（openpyxl）。每个人脸一行。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import APP_NAME, VERSION
from .pipeline import ST_NO_FACE, ImageResult
from .interval import Calibration

HEADERS = [
    "原始文件名",
    "图片完整路径",
    "图片中的人脸编号",
    "预测年龄",
    "年龄区间下限",
    "年龄区间上限",
    "年龄区间类型",
    "使用的模型",
    "处理状态",
    "错误信息",
    "区间半宽(±岁)",
    "检测置信度",
    "多视图标准差",
    "处理耗时(秒)",
    "图片宽",
    "图片高",
    "对比模型",
    "对比年龄(岁)",
]

WIDTHS = [22, 52, 14, 10, 12, 12, 20, 34, 16, 40, 12, 12, 12, 12, 9, 9, 34, 12]

_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def export_excel(
    results: list[ImageResult],
    out_path: str | Path,
    model_name: str,
    mode_label: str,
    half_width: int,
    calib: Calibration | None = None,
    notes: str = "",
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "年龄推断结果"

    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[1].height = 30

    for r in results:
        fname = Path(r.path).name
        if not r.faces:
            ws.append([fname, r.path, "-", "", "", "", "", model_name, r.status, r.error,
                       "", "", "", round(r.elapsed, 3), r.width, r.height, "", ""])
            continue
        for f in r.faces:
            ws.append([
                fname,
                r.path,
                f.index,
                None if f.age is None else round(float(f.age), 1),
                f.lo if f.age is not None else "",
                f.hi if f.age is not None else "",
                f.itype,
                model_name,
                f.status if r.status == "成功" else f"{f.status}",
                f.error,
                round(f.half, 1),
                round(f.det_score, 3),
                "" if f.sigma is None else round(f.sigma, 3),
                round(r.elapsed, 3),
                r.width,
                r.height,
                f.model2,
                None if f.age2 is None else round(float(f.age2), 1),
            ])

    last_row = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last_row, max_col=len(HEADERS)):
        for cell in row:
            cell.border = _BORDER
            cell.font = Font(size=10)
        row[0].alignment = Alignment(vertical="center")
        row[1].alignment = Alignment(vertical="center")
        for i in (2, 3, 4, 5, 10, 11, 12, 13, 14, 15, 16, 17):
            row[i].alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    if last_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"

    # ---------------- 说明 sheet ---------------- #
    ws2 = wb.create_sheet("说明")
    compare_models = sorted({f.model2 for r in results for f in r.faces if f.model2})
    info = [
        ("软件", f"{APP_NAME} v{VERSION}"),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("年龄模型", model_name),
        ("对比模型", "、".join(compare_models) if compare_models else "未启用"),
        ("区间模式", mode_label),
        ("基础区间半宽", f"±{half_width} 岁"),
        ("", ""),
        ("重要说明 1", "基础模式的年龄区间 = 模型预测值 ± 指定宽度，属于人为设定的参考范围，"),
        ("", "不是统计置信区间，不代表真实年龄一定落在其中。"),
        ("重要说明 2", "若使用校准区间（保形预测），区间由带真实年龄标签的公开验证集标定，"),
        ("", "其覆盖率是在该验证集上实测的数值，对分布差异较大的照片不保证同样效果。"),
        ("重要说明 3", "所有年龄均为模型基于人脸外观的推断（表观年龄），未读取图片中的任何文字标签。"),
        ("重要说明 4", "模型只输出年龄，不推断性别、种族等属性；照片不会上传到任何服务器。"),
        ("", ""),
        ("图片总数", len(results)),
        ("人脸总数", sum(len(r.faces) for r in results)),
        ("处理成功", sum(1 for r in results if r.status == "成功")),
        ("未检测到人脸", sum(1 for r in results if r.status == ST_NO_FACE)),
        ("失败", sum(1 for r in results if r.status in ("读取失败", "处理失败", "部分人脸失败"))),
    ]
    if calib is not None:
        info += [
            ("", ""),
            ("校准信息", calib.describe()),
            ("校准数据集", calib.dataset),
            ("校准样本数", calib.n_calib),
            ("测试样本数", calib.n_test),
            ("校准模型", calib.model),
            ("校准时间", calib.created),
        ]
    if notes:
        info += [("", ""), ("备注", notes)]

    for k, v in info:
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 90
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=2):
        row[0].font = Font(bold=True, size=10)
        row[1].font = Font(size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="center")

    # ---------------- 统计 sheet ---------------- #
    ws3 = wb.create_sheet("分布统计")
    ws3.append(["年龄区间(预测值)", "人脸数"])
    buckets = [(0, 12), (13, 19), (20, 29), (30, 39), (40, 49), (50, 59), (60, 69), (70, 200)]
    ages = [f.age for r in results for f in r.faces if f.age is not None]
    for lo, hi in buckets:
        n = sum(1 for a in ages if lo <= a <= hi)
        label = f"{lo}-{hi}" if hi < 200 else f"{lo}+"
        ws3.append([label, n])
    ws3.append(["合计", len(ages)])
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 10
    for c in ws3[1]:
        c.fill = _HEAD_FILL
        c.font = _HEAD_FONT

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
